import csv
import io
import json
import re
from typing import Annotated, Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile

from auth.router import get_current_user
from auth.schemas import User
from database.schema import SchemaService
from database.service import DatabaseService
from datasource.schema_enrichment import SchemaEnrichmentService
from datasource.schemas import DataSourceCreate, DataSourceUpdate, GrantRevokeRequest
from datasource.service import DataSourceService
from datasource.suggestion import SuggestionService
from dependencies import get_db_service, get_schema_service
from llm.service import LlmService

router = APIRouter(prefix="/datasources", tags=["datasources"])

MAX_FILE_ROWS = 50_000
MAX_FILE_SIZE = 20 * 1024 * 1024

NUMERIC_RE = re.compile(r"^-?(\d+\.?\d*|\d*\.\d+)$")
THOUSAND_SEP_RE = re.compile(r"^-?\d{1,3}(,\d{3})+(\.\d+)?$")
INVISIBLE_RE = re.compile(r"[\u200B\uFEFF\u00A0]")
COL_SANITIZE_RE = re.compile(r"[^a-zA-Z0-9_\u4e00-\u9fa5]")


def _normalize_value(v: Any) -> Any:
    """清除零宽字符、千分位、空串 -> None，保持类型推断稳定。"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        # NaN/inf -> None
        if isinstance(v, float) and not (v == v) or v in (float("inf"), float("-inf")):
            return None
        return v
    if isinstance(v, str):
        s = INVISIBLE_RE.sub(" ", v).strip()
        if s == "":
            return None
        if THOUSAND_SEP_RE.match(s):
            s = s.replace(",", "")
        return s
    return v


def _is_numeric(v: Any) -> bool:
    if isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return v == v and v not in (float("inf"), float("-inf"))
    if isinstance(v, str):
        return bool(NUMERIC_RE.match(v))
    return False


def _infer_mysql_type(values: list[Any]) -> str:
    non_null = [v for v in values if v is not None]
    if not non_null:
        return "TEXT"
    if not all(_is_numeric(v) for v in non_null):
        return "TEXT"
    has_decimal = any(
        (isinstance(v, float) and not float(v).is_integer())
        or (isinstance(v, str) and "." in v)
        for v in non_null
    )
    return "DOUBLE" if has_decimal else "BIGINT"


def _coerce(value: Any, type_: str) -> Any:
    if value is None:
        return None
    if type_ == "TEXT":
        return value if isinstance(value, str) else str(value)
    if isinstance(value, (int, float)):
        return value
    try:
        n = float(value)
        if n != n or n in (float("inf"), float("-inf")):
            return None
        return n
    except (ValueError, TypeError):
        return None


def _ensure_unique_column_names(raw_names: list[str]) -> list[str]:
    def sanitize(s: str) -> str:
        s = COL_SANITIZE_RE.sub("_", s or "")
        s = s.strip("_")
        return s or "col"

    seen: dict[str, int] = {}
    result: list[str] = []
    for raw in raw_names:
        name = sanitize(raw)
        count = seen.get(name, 0)
        seen[name] = count + 1
        result.append(name if count == 0 else f"{name}_{count}")
    return result


async def get_llm_service() -> LlmService:
    return LlmService()


async def get_enrichment_service(
    db: Annotated[DatabaseService, Depends(get_db_service)],
    llm: Annotated[LlmService, Depends(get_llm_service)],
) -> SchemaEnrichmentService:
    return SchemaEnrichmentService(db, llm)


async def get_datasource_service(
    db: Annotated[DatabaseService, Depends(get_db_service)],
    schema: Annotated[SchemaService, Depends(get_schema_service)],
    enrichment: Annotated[
        SchemaEnrichmentService, Depends(get_enrichment_service)
    ],
) -> DataSourceService:
    return DataSourceService(db, schema, enrichment)


async def get_suggestion_service(
    db: Annotated[DatabaseService, Depends(get_db_service)],
    schema: Annotated[SchemaService, Depends(get_schema_service)],
    llm: Annotated[LlmService, Depends(get_llm_service)],
) -> SuggestionService:
    return SuggestionService(db, schema, llm)


@router.get("")
async def find_all(
    user: Annotated[User, Depends(get_current_user)],
    service: Annotated[DataSourceService, Depends(get_datasource_service)],
):
    return await service.find_all_for_user(user.id)


@router.post("")
async def create(
    body: DataSourceCreate,
    user: Annotated[User, Depends(get_current_user)],
    service: Annotated[DataSourceService, Depends(get_datasource_service)],
):
    if not body.name or not body.name.strip():
        raise HTTPException(status_code=400, detail="请提供数据源名称")
    return await service.create(user.id, body.name.strip(), body.description)


@router.patch("/{id}")
async def update(
    id: int,
    body: DataSourceUpdate,
    user: Annotated[User, Depends(get_current_user)],
    service: Annotated[DataSourceService, Depends(get_datasource_service)],
):
    if body.name is not None and not body.name.strip():
        raise HTTPException(status_code=400, detail="数据源名称不能为空")
    return await service.update(
        id, user.id, body.name.strip() if body.name else None, body.description
    )


@router.delete("/{id}")
async def remove(
    id: int,
    user: Annotated[User, Depends(get_current_user)],
    service: Annotated[DataSourceService, Depends(get_datasource_service)],
):
    await service.remove(id, user.id)
    return {"success": True}


@router.get("/{id}/tables")
async def list_tables(
    id: int,
    user: Annotated[User, Depends(get_current_user)],
    service: Annotated[DataSourceService, Depends(get_datasource_service)],
):
    return await service.list_uploaded_tables(id, user.id)


@router.post("/{id}/tables")
async def upload_table(
    id: int,
    user: Annotated[User, Depends(get_current_user)],
    service: Annotated[DataSourceService, Depends(get_datasource_service)],
    suggestion: Annotated[SuggestionService, Depends(get_suggestion_service)],
    file: UploadFile = File(...),
    tableConfigs: str = Form(...),
):
    """CSV/Excel -> 数据清洗 -> 类型推断 -> CREATE + INSERT；多 sheet 支持。"""
    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=400, detail="文件过大，上限 20MB")

    try:
        configs = json.loads(tableConfigs)
    except Exception:
        raise HTTPException(status_code=400, detail="表配置格式错误")
    if not configs or not isinstance(configs, list):
        raise HTTPException(status_code=400, detail="请选择至少一个要上传的表")

    ext = (file.filename or "").rsplit(".", 1)[-1].lower()
    workbook = None
    if ext in ("xlsx", "xls"):
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=True
            )
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Excel 解析失败: {e}")

    results = []
    for config in configs:
        display_name = (config.get("displayName") or "").strip()
        if not display_name:
            continue

        raw_rows: list[dict[str, Any]] = []

        if ext == "csv":
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = content.decode("gbk", errors="replace")
            reader = csv.DictReader(io.StringIO(text))
            for row in reader:
                raw_rows.append(dict(row))
                if len(raw_rows) > MAX_FILE_ROWS:
                    raise HTTPException(
                        status_code=400,
                        detail=f"文件行数超过上限 {MAX_FILE_ROWS} 行",
                    )
        elif workbook is not None:
            sheet_name = config.get("sheetName") or workbook.sheetnames[0]
            if sheet_name not in workbook.sheetnames:
                continue
            sheet = workbook[sheet_name]
            rows_iter = sheet.iter_rows(values_only=True)
            try:
                header = list(next(rows_iter))
            except StopIteration:
                continue
            header = [str(h) if h is not None else "" for h in header]
            for row in rows_iter:
                raw_rows.append(
                    {header[i]: row[i] for i in range(min(len(header), len(row)))}
                )
                if len(raw_rows) > MAX_FILE_ROWS:
                    raise HTTPException(
                        status_code=400,
                        detail=f"Sheet「{sheet_name}」行数超过上限 {MAX_FILE_ROWS} 行",
                    )
        else:
            continue

        if not raw_rows:
            continue

        col_names = list(raw_rows[0].keys())
        normalized_rows = [
            {k: _normalize_value(row.get(k)) for k in col_names}
            for row in raw_rows
        ]

        unique_names = _ensure_unique_column_names(col_names)
        columns = [
            {
                "name": unique_names[i],
                "originalName": col_names[i],
                "type": _infer_mysql_type(
                    [r[col_names[i]] for r in normalized_rows]
                ),
            }
            for i in range(len(col_names))
        ]

        clean_rows = []
        for row in normalized_rows:
            cleaned: dict[str, Any] = {}
            for i, orig in enumerate(col_names):
                cleaned[columns[i]["name"]] = _coerce(
                    row.get(orig), columns[i]["type"]
                )
            clean_rows.append(cleaned)

        result = await service.upload_table(
            id, user.id, display_name, columns, clean_rows
        )
        results.append(result)

    if results:
        suggestion.invalidate_and_regenerate(id, user.id)

    return results


@router.delete("/{id}/tables/{table_id}")
async def delete_table(
    id: int,
    table_id: int,
    user: Annotated[User, Depends(get_current_user)],
    service: Annotated[DataSourceService, Depends(get_datasource_service)],
    suggestion: Annotated[SuggestionService, Depends(get_suggestion_service)],
):
    ds_id = await service.delete_uploaded_table(table_id, user.id)
    if ds_id:
        suggestion.invalidate_and_regenerate(ds_id, user.id)
    return {"ok": True}


@router.get("/{id}/suggestions")
async def get_suggestions(
    id: int,
    user: Annotated[User, Depends(get_current_user)],
    service: Annotated[SuggestionService, Depends(get_suggestion_service)],
):
    questions = await service.get_suggestions(id, user.id)
    return {"questions": questions}


@router.get("/{id}/schema")
async def get_schema(
    id: int,
    user: Annotated[User, Depends(get_current_user)],
    schema: Annotated[SchemaService, Depends(get_schema_service)],
):
    return await schema.get_structured_schema(id, user.id)


@router.get("/{id}")
async def find_one(
    id: int,
    user: Annotated[User, Depends(get_current_user)],
    service: Annotated[DataSourceService, Depends(get_datasource_service)],
):
    return await service.find_one(id, user.id)


@router.post("/{id}/grant")
async def grant(
    id: int,
    body: GrantRevokeRequest,
    user: Annotated[User, Depends(get_current_user)],
    service: Annotated[DataSourceService, Depends(get_datasource_service)],
):
    if not body.userId:
        raise HTTPException(status_code=400, detail="请提供 userId")
    await service.grant(id, body.userId, user.id)
    return {"ok": True}


@router.post("/{id}/revoke")
async def revoke(
    id: int,
    body: GrantRevokeRequest,
    user: Annotated[User, Depends(get_current_user)],
    service: Annotated[DataSourceService, Depends(get_datasource_service)],
):
    if not body.userId:
        raise HTTPException(status_code=400, detail="请提供 userId")
    await service.revoke(id, body.userId, user.id)
    return {"ok": True}


@router.get("/{id}/permissions")
async def list_permissions(
    id: int,
    user: Annotated[User, Depends(get_current_user)],
    service: Annotated[DataSourceService, Depends(get_datasource_service)],
):
    return await service.list_permission_users(id, user.id)
